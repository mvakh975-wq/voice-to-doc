from docx import Document

from voice_to_doc.render import to_docx, to_markdown

LECTURE = {"title": "История сети", "summary": ["Тезис"], "kind": "lecture", "duration_s": 600.0, "warnings": [],
           "items": [{"title": "Введение", "start": 0.0, "paragraphs": ["Абзац 1.", "Абзац 2."]},
                     {"title": "ARPANET", "start": 452.0, "paragraphs": ["Текст."]}]}
INTERVIEW = {"title": "Беседа", "summary": [], "kind": "interview", "duration_s": 4000.0,
             "warnings": ["Фрагмент с 1:00:00 не оформлен — показан исходный текст распознавания."],
             "items": [{"speaker": "Интервьюер", "start": 190.0, "text": "Вопрос?"},
                       {"speaker": None, "start": 3600.0, "text": "сырой текст", "raw": True}]}


def test_lecture_markdown():
    assert to_markdown(LECTURE) == (
        "# История сети\n\n## Кратко\n- Тезис\n\n"
        "## [00:00] Введение\n\nАбзац 1.\n\nАбзац 2.\n\n"
        "## [07:32] ARPANET\n\nТекст.\n"
    )


def test_interview_markdown_long_timestamps_and_warnings():
    assert to_markdown(INTERVIEW) == (
        "# Беседа\n\n"
        "**Интервьюер [0:03:10]:** Вопрос?\n\n"
        "**Не определено [1:00:00]:** сырой текст\n\n"
        "---\n\n**Замечания:**\n- Фрагмент с 1:00:00 не оформлен — показан исходный текст распознавания.\n"
    )


def test_docx_structure(tmp_path):
    path = tmp_path / "r.docx"
    to_docx(LECTURE, path)
    paras = [(p.style.name, p.text) for p in Document(path).paragraphs]
    assert paras == [("Heading 1", "История сети"), ("Heading 2", "Кратко"), ("List Bullet", "Тезис"),
                     ("Heading 2", "[00:00] Введение"), ("Normal", "Абзац 1."), ("Normal", "Абзац 2."),
                     ("Heading 2", "[07:32] ARPANET"), ("Normal", "Текст.")]


def test_docx_interview_bold_speaker(tmp_path):
    path = tmp_path / "r.docx"
    to_docx(INTERVIEW, path)
    paras = Document(path).paragraphs
    first = paras[1]
    assert first.text == "Интервьюер [0:03:10]: Вопрос?"
    assert first.runs[0].bold and not first.runs[1].bold
    assert paras[-1].style.name == "List Bullet" and "не оформлен" in paras[-1].text
