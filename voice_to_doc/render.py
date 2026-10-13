"""Render a structured document to Markdown and Word."""

from pathlib import Path

from docx import Document

from voice_to_doc.structure import fmt_ts

UNKNOWN_SPEAKER = "Не определено"


def _ts(doc: dict, seconds: float) -> str:
    return fmt_ts(seconds, doc["duration_s"] >= 3600)


def to_markdown(doc: dict) -> str:
    out = [f"# {doc['title']}", ""]
    if doc["summary"]:
        out += ["## Кратко", *[f"- {s}" for s in doc["summary"]], ""]
    for it in doc["items"]:
        if doc["kind"] == "lecture":
            out += [f"## [{_ts(doc, it['start'])}] {it['title']}", ""]
            for p in it["paragraphs"]:
                out += [p, ""]
        else:
            out += [f"**{it['speaker'] or UNKNOWN_SPEAKER} [{_ts(doc, it['start'])}]:** {it['text']}", ""]
    if doc["warnings"]:
        out += ["---", "", "**Замечания:**", *[f"- {w}" for w in doc["warnings"]], ""]
    return "\n".join(out)


def to_docx(doc: dict, path: Path) -> None:
    d = Document()
    d.add_heading(doc["title"], level=1)
    if doc["summary"]:
        d.add_heading("Кратко", level=2)
        for s in doc["summary"]:
            d.add_paragraph(s, style="List Bullet")
    for it in doc["items"]:
        if doc["kind"] == "lecture":
            d.add_heading(f"[{_ts(doc, it['start'])}] {it['title']}", level=2)
            for p in it["paragraphs"]:
                d.add_paragraph(p)
        else:
            para = d.add_paragraph()
            para.add_run(f"{it['speaker'] or UNKNOWN_SPEAKER} [{_ts(doc, it['start'])}]: ").bold = True
            para.add_run(it["text"])
    if doc["warnings"]:
        d.add_heading("Замечания", level=2)
        for w in doc["warnings"]:
            d.add_paragraph(w, style="List Bullet")
    d.save(path)
